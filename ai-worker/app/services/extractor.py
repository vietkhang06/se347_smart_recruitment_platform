"""Safe CV text extraction engine supporting PDF and DOCX formats."""

import io
from typing import Tuple

from pypdf import PdfReader
from docx import Document


MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB
MAX_PDF_PAGES = 20
MAX_CHARACTERS = 500_000


def sanitize_text(text: str) -> str:
    """Sanitize extracted text, stripping null bytes and bounding length."""
    if not text:
        return ""
    # Strip null bytes
    sanitized = text.replace("\x00", "")
    # Normalize excessive newlines
    lines = [line.strip() for line in sanitized.splitlines()]
    clean_text = "\n".join(line for line in lines if line)
    return clean_text[:MAX_CHARACTERS]


def extract_cv_text(filename: str, content: bytes, content_type: str = "") -> Tuple[str, int]:
    """
    Extract text and page count from binary PDF or DOCX content.
    Returns (raw_text, page_count).
    Raises ValueError for invalid, oversized, or corrupted files.
    """
    if len(content) > MAX_FILE_SIZE_BYTES:
        raise ValueError("File size exceeds 10MB limit.")

    if not content:
        raise ValueError("File content is empty.")

    lower_filename = filename.lower()
    is_pdf = lower_filename.endswith(".pdf") or "pdf" in content_type.lower()
    is_docx = lower_filename.endswith(".docx") or "wordprocessingml" in content_type.lower()

    if is_pdf:
        # Validate PDF magic bytes
        if b"%PDF" not in content[:1024]:
            raise ValueError("Corrupted or invalid PDF header.")

        try:
            reader = PdfReader(io.BytesIO(content))
            page_count = len(reader.pages)
            if page_count > MAX_PDF_PAGES:
                raise ValueError(f"PDF exceeds maximum {MAX_PDF_PAGES} pages limit (found {page_count}).")

            extracted_pages = []
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    extracted_pages.append(text)

            raw_text = sanitize_text("\n\n".join(extracted_pages))
            if not raw_text.strip():
                raw_text = "CV Document (Non-text or scanned PDF)"
            return raw_text, page_count
        except Exception as e:
            if "pages limit" in str(e):
                raise
            raise ValueError(f"Failed to parse PDF document: {str(e)}")

    elif is_docx:
        # Validate DOCX PK zip magic bytes
        if not content.startswith(b"PK\x03\x04"):
            raise ValueError("Corrupted or invalid DOCX header.")

        try:
            doc = Document(io.BytesIO(content))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

            # Extract from tables as well
            for table in doc.tables:
                for row in table.rows:
                    row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_texts:
                        paragraphs.append(" | ".join(row_texts))

            raw_text = sanitize_text("\n".join(paragraphs))
            if not raw_text.strip():
                raw_text = "CV Document (Empty Word document)"
            return raw_text, 1
        except Exception as e:
            raise ValueError(f"Failed to parse DOCX document: {str(e)}")

    else:
        raise ValueError("Unsupported file format. Only PDF (.pdf) and Word (.docx) files are supported.")
