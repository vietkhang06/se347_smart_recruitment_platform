"""Tests for Phase 3 CV extraction, vector embeddings, and match scoring."""

import io
import math
import pytest
from fastapi.testclient import TestClient
from docx import Document
from pypdf import PdfWriter

from app.main import app


@pytest.fixture
def client():
    with TestClient(app) as test_client:
        yield test_client


def create_sample_pdf() -> bytes:
    """Generate in-memory sample PDF bytes."""
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    # Write empty page or text
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def create_sample_docx() -> bytes:
    """Generate in-memory sample DOCX bytes with resume content."""
    doc = Document()
    doc.add_heading("Nguyen Van A", level=1)
    doc.add_paragraph("Senior Java Developer with 5 years experience in Spring Boot, PostgreSQL, and Docker.")
    doc.add_paragraph("Email: nguyenvana@example.com | Phone: 0901234567")
    doc.add_paragraph("Education: University of Science, B.S. Computer Science")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_cv_extract_docx(client: TestClient):
    docx_bytes = create_sample_docx()
    response = client.post(
        "/api/v1/cv/extract",
        files={"file": ("resume.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 200
    data = response.json()
    assert "Nguyen Van A" in data["raw_text"]
    assert data["parsed_data"]["email"] == "nguyenvana@example.com"
    assert data["parsed_data"]["phone"] == "0901234567"
    assert "Java" in data["parsed_data"]["skills"]
    assert "Spring Boot" in data["parsed_data"]["skills"]
    assert data["char_count"] > 50


def test_cv_extract_pdf(client: TestClient):
    pdf_bytes = create_sample_pdf()
    response = client.post(
        "/api/v1/cv/extract",
        files={"file": ("cv.pdf", pdf_bytes, "application/pdf")},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["page_count"] == 1
    assert "raw_text" in data
    assert "parsed_data" in data


def test_cv_extract_unsupported_format(client: TestClient):
    response = client.post(
        "/api/v1/cv/extract",
        files={"file": ("resume.txt", b"plain text content", "text/plain")},
    )
    assert response.status_code == 400
    assert "Unsupported file format" in response.json()["detail"]


def test_cv_extract_corrupt_pdf(client: TestClient):
    response = client.post(
        "/api/v1/cv/extract",
        files={"file": ("corrupt.pdf", b"NOT_A_REAL_PDF_DATA", "application/pdf")},
    )
    assert response.status_code == 400
    assert "Corrupted or invalid PDF header" in response.json()["detail"]


def test_embeddings_generation_deterministic(client: TestClient):
    payload = {
        "texts": [
            "Senior Backend Engineer with Java and Spring Boot experience",
            "Seeking Frontend Developer proficient in React and TypeScript"
        ]
    }
    response1 = client.post("/api/v1/embeddings", json=payload)
    assert response1.status_code == 200
    data1 = response1.json()
    assert data1["dimension"] == 1536
    assert len(data1["embeddings"]) == 2
    assert len(data1["embeddings"][0]) == 1536

    # Verify unit Euclidean norm (L2 norm ≈ 1.0)
    v1 = data1["embeddings"][0]
    l2_norm = math.sqrt(sum(x * x for x in v1))
    assert pytest.approx(l2_norm, 0.01) == 1.0

    # Repeat call and verify identical output (deterministic)
    response2 = client.post("/api/v1/embeddings", json=payload)
    assert response2.status_code == 200
    data2 = response2.json()
    assert data1["embeddings"] == data2["embeddings"]


def test_match_score_calculation(client: TestClient):
    payload = {
        "job_title": "Senior Backend Developer",
        "job_description": "We need an engineer to build scalable microservices.",
        "job_requirements": "Must have Java, Spring Boot, PostgreSQL, Docker.",
        "cv_text": "Nguyen Van A. 5 years experience with Java, Spring Boot, and PostgreSQL. Docker proficiency.",
        "candidate_skills": ["Java", "Spring Boot", "PostgreSQL", "Docker"]
    }
    response = client.post("/api/v1/match/score", json=payload)
    assert response.status_code == 200
    data = response.json()
    assert 0 <= data["overall_score"] <= 100
    assert "required_skills_matched" in data["breakdown"]
    assert "Java" in data["breakdown"]["required_skills_matched"]
    assert data["rationale"] is not None
